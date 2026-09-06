"""Build an editable, minimal PowerPoint; preserve rehearsal notes from slides.md."""
from pathlib import Path
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'SyFI_ML_Serving_refined.pptx'
WHITE = 'FFFFFF'
INK = '18283B'
BLUE = '2463A6'
MUTED = '596777'
PALE = 'EDF3FA'
GRAY = 'F5F7FA'

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = 'Serving frontier MoE models on 8 H100s'
prs.core_properties.author = 'Tan Ngo'
prs.core_properties.subject = 'ML systems experiments for Professor Kan Zhu and UW SyFI'
notes = re.findall(r'<!--(.*?)-->', (ROOT / 'slides.md').read_text(encoding='utf-8'), re.S)
assert len(notes) == 20


def text(slide, value, x, y, w, h, size=24, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Aptos'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(12)
        if align is not None:
            p.alignment = align
    return box


def panel(slide, value, y=5.77, size=24):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.6), Inches(y), Inches(12.1), Inches(.72))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(PALE)
    shape.line.fill.background()
    text(slide, value, .82, y + .14, 11.65, .46, size, True, BLUE)


def base(title, source='', caveat=''):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    i = len(prs.slides)
    text(s, 'UW SyFI  /  ML SYSTEMS', .6, .28, 10, .25, 11, True, BLUE)
    text(s, title, .6, .88, 12.05, 1.7 if '\n' in title else .94, 31, True)
    if caveat:
        text(s, caveat, .63, 6.57, 12, .43, 14, False, MUTED)
    text(s, source or 'Tan Ngo · 5 September 2026', .63, 7.08, 11.45, .2, 10, False, MUTED)
    text(s, str(i) if i <= 12 else f'B{i-12}', 12.1, 7.03, .6, .3, 12, False, MUTED, PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes[i-1].strip()
    if i > 12:
        s._element.set('show', '0')
    return s


def table(slide, headers, rows, widths, y=2.0, row_h=.64, size=23, highlight=None):
    shape = slide.shapes.add_table(len(rows)+1, len(headers), Inches(.65), Inches(y), Inches(12), Inches(row_h*(len(rows)+1)))
    t = shape.table
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for r, row in enumerate([headers] + rows):
        for c, value in enumerate(row):
            cell = t.cell(r, c)
            cell.text = str(value)
            cell.margin_left = Inches(.16)
            cell.margin_right = Inches(.12)
            cell.margin_top = Inches(.08)
            cell.margin_bottom = Inches(.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(BLUE if r == 0 else PALE if r == highlight else WHITE if r % 2 else GRAY)
            for p in cell.text_frame.paragraphs:
                p.font.name = 'Aptos'
                p.font.size = Pt(size if r else size-2)
                p.font.bold = r == 0 or r == highlight
                p.font.color.rgb = RGBColor.from_string(WHITE if r == 0 else INK)
                p.space_after = Pt(0)
    return t


def bullets(slide, items, y=2.05, step=.88, size=26):
    for i, item in enumerate(items):
        text(slide, '•', .7, y+i*step, .25, .5, size, True, BLUE)
        text(slide, item, 1.1, y+i*step, 11.2, step-.12, size)


s = base('Serving frontier MoE models\non 8 H100 GPUs')
text(s, 'Concurrency · Context · Prefix caching · Speculation', .65, 3.0, 12, .65, 27, False, MUTED)
text(s, 'Tan Ngo', .65, 4.25, 10, .6, 27, True)
text(s, 'Professor Kan Zhu and UW SyFI\n5 September 2026', .65, 4.95, 11, .9, 21, False, MUTED)
panel(s, 'Three measured models. Four systems questions.', y=6.02)

s = base('Four questions guide the study', 'Source: task.md; report.md §§1–2')
bullets(s, ['What limits performance as the workload changes?', 'What makes prefix caching difficult?', 'When does speculative decoding help?', 'What does serving cost?'])
panel(s, 'GLM, DeepSeek, Qwen measured; Kimi architecture only.')

s = base('Same MoE idea. Different attention state.', 'Source: report.md §2; model configurations', 'Architecture facts; these do not establish the cause of measured speed differences.')
table(s, ['Model', 'Attention layers', 'Experts / selected'], [
    ['GLM-5.3-Flash', '34 recurrent + 11 sparse', '288 / 8'],
    ['DeepSeek-V4-Flash', '43 sparse; compressed history', '256 / 6'],
    ['Qwen3.8-Flash-Next', '36 recurrent + 12 QSA', '512 / 10'],
    ['Kimi K3 (unmeasured)', '69 recurrent + 24 gated MLA', '896 / 16'],
], [3.5, 5.3, 3.2], row_h=.67, size=21)
panel(s, 'Recurrent state stays fixed; retained history grows.', size=24)

s = base('A fixed GPU budget; three workload axes', 'Source: bench.sh; report.md §1', 'Synthetic prompts; mostly one retained run per point; builds and backends differ.')
text(s, '8 × H100 80GB  ·  Tensor + expert parallelism', .65, 1.93, 12, .6, 25, True, BLUE)
table(s, ['Axis', 'Input / output', 'Concurrency'], [
    ['Batch', '16K / 256 tokens', '1–64'],
    ['Context', '16K–260K / 256 tokens', '8'],
    ['Prefix sharing', '64K shared + 2K unique / 256', '8'],
], [2.4, 6.5, 3.1], y=2.75, row_h=.62, size=23)
panel(s, 'Output tok/s includes input work; TTFT includes queueing.', size=22)

s = base('Qwen leads the corrected 16K throughput grid', 'Source: report.md §3; main batch arms', 'Output tokens/s, MTP off. Deployment comparison; response quality was not measured.')
table(s, ['Concurrency', 'Qwen', 'GLM', 'DeepSeek'], [
    ['1', '106.2', '96.3', '85.0'], ['4', '256.8', '225.3', '219.8'],
    ['16', '420.2', '359.2', '330.9'], ['64', '517.7', '447.1', '389.4'],
], [3.0, 3.0, 3.0, 3.0], row_h=.65, size=25, highlight=4)
panel(s, 'At concurrency 64: 1.16× GLM and 1.33× DeepSeek.')

s = base('More throughput has a latency price', 'Source: report.md §3; GLM bf16kv; DeepSeek util085-dev20073', 'DeepSeek uses a separate build/utilization arm. These are not universal optimal settings.')
text(s, 'Increasing concurrency from 8 to 64', .65, 1.91, 12, .5, 25, False, MUTED)
table(s, ['Model', 'Throughput increase', 'Token-latency increase'], [
    ['GLM', '1.51×', '7.15×'], ['DeepSeek', '1.37×', '9.41×'],
], [3.0, 4.5, 4.5], y=2.72, row_h=.82, size=26)
panel(s, 'Choose concurrency against a latency objective.')

s = base('Long input changes which metric wins', 'Source: report.md §3; ctx_isl131072_c8.json', 'Qwen context uses an earlier, smaller cache pool; a controlled rerun is needed.')
text(s, '131K input  ·  256 output  ·  Concurrency 8', .65, 1.9, 12, .5, 24, False, MUTED)
table(s, ['Model', 'Output tokens/s', 'First-token latency'], [
    ['Qwen', '56.5', '19.84 s'], ['GLM', '47.2', '12.89 s'], ['DeepSeek', '34.7', '26.66 s'],
], [4, 4, 4], y=2.65, row_h=.65, size=25)
panel(s, 'Highest throughput ≠ fastest first token.')

s = base('Prefix reuse needs consistent model state', 'Source: report.md §4; prefix_p64k_n*.json', '64 requests; 64K shared + 2K unique input; concurrency 8. No cache-on/off control.')
table(s, ['Distinct prefixes', 'GLM tokens/s', 'DeepSeek tokens/s'], [
    ['1', '265.3', '198.2'], ['4', '365.9', '391.6'], ['16', '199.4', '114.6'],
], [4, 4, 4], row_h=.68, size=25, highlight=2)
text(s, 'Maximum sharing did not maximize throughput.', .7, 5.05, 12, .52, 25)
panel(s, 'Restore both KV and recurrent state at one boundary.', size=23)

s = base('MTP helps some workloads, hurts others', 'Source: report.md §5; controlled GLM base / MTP arms', 'Ratios are throughput with MTP ÷ base throughput. Scheduling explanations remain hypotheses.')
table(s, ['GLM concurrency', '1 draft token', '5 draft tokens'], [
    ['1', '1.24×', '1.25×'], ['4', '1.06×', '1.02×'], ['16', '0.96×', '0.95×'], ['64', '0.99×', '0.91×'],
], [4, 4, 4], row_h=.65, size=25)
panel(s, 'At 131K: first-token latency −17%; throughput −4%.', size=23)

s = base('Serving cost depends on the workload', 'Source: report.md §6; calculated from raw throughput', 'Includes input work. Illustrative price, not a quote; no quality adjustment or idle-time cost.')
text(s, 'Assume $2.50/GPU-hour → $20/node-hour', .65, 1.92, 12, .55, 27, True, BLUE)
text(s, '16K input / 256 output  ·  Concurrency 64', .65, 2.6, 12, .5, 23, False, MUTED)
table(s, ['Qwen', 'GLM', 'DeepSeek'], [['$10.73', '$12.43', '$14.27']], [4, 4, 4], y=3.4, row_h=.83, size=34)
panel(s, 'Dollars per million output tokens', size=25)

s = base('Debugging connects architecture to the result', 'Source: fix_bug.md Bugs 1, 6, 12; report.md §2')
table(s, ['Feature', 'Observed failure', 'Lesson'], [
    ['Recurrent layers', 'Sequence-state capacity\nblocks startup', 'Budget state per\nactive sequence'],
    ['Prefix caching', 'Warmup creates\n16,000 reused tokens', 'Warm kernels; use\nfresh test prefixes'],
    ['Automatic pool sizing', 'Qwen c4 changes\n162.2 → 256.8 tok/s', 'Validate startup before\nexplaining speed'],
], [3.7, 4.1, 4.2], row_h=.88, size=21)
panel(s, 'Symptom → discriminating test → fix → bounded conclusion', size=22)

s = base('Next: validate, profile, then improve goodput', 'Source: report.md §8')
bullets(s, ['Validate: matched startup and repeated runs.', 'Profile: separate prefill, decode, and communication.', 'Improve: state-aware caching and scheduling.'], step=1.03, size=27)
panel(s, 'Serving choices depend on workload, state, and runtime.', size=23)
text(s, 'Questions', .65, 6.64, 11, .38, 22, True, BLUE)

s = base('Backup · Is the comparison fair?', 'Source: report.md §1')
table(s, ['Controlled', 'Different or missing'], [
    ['GPU class and count', 'Physical topology not established'],
    ['Workload target lengths', 'Tokenizers and task quality'],
    ['MTP off in main comparison', 'Precision, builds, backends, some pools'],
], [5, 7], row_h=.85, size=23)
panel(s, 'Deployment comparison, not isolated architecture.', size=24)

s = base('Backup · What is the actual bottleneck?', 'Source: report.md §3')
table(s, ['Candidate', 'Evidence needed'], [
    ['GEMMs / expert routing', 'Operator timing and per-rank load'],
    ['Communication / launches', 'Timeline and TP/EP sweep'],
    ['HBM traffic', 'Hardware counters within decode'],
    ['Scheduling / cache pressure', 'Chunk-size A/B and preemptions'],
], [5, 7], row_h=.69, size=23)
panel(s, 'The current data identifies regimes, not a dominant kernel.', size=22)

s = base('Backup · Can recurrent state be cached?', 'Source: report.md §4')
bullets(s, ['Yes: save state at compatible prefix boundaries.', 'KV and recurrent checkpoints must agree on position.', 'Branches and rollback need safe copies or restoration.', 'Checkpoint granularity trades storage for recomputation.'], size=25)
panel(s, 'Recurrent state is not inherently uncacheable.')

s = base('Backup · Why not compare parameter counts?', 'Source: report.md §2; prior tensor analyses and Kimi official card')
bullets(s, ['Resident weights differ from per-token selected weights.', 'Batched tokens can select different experts.', 'Qwen has a 51.23B lookup table, not a full per-token GEMM.', 'State, indexing, and communication add other costs.'], size=24)
panel(s, 'Kimi K3: ideal 4-bit weights alone ≈ 1,304 GiB.', size=24)

s = base('Backup · Enable MTP or FP8 KV?', 'Source: report.md §§5, 7')
table(s, ['GLM option', 'Observed tradeoff'], [
    ['MTP, 1 draft', '+24% throughput at c1; approximately flat at c64'],
    ['MTP, 5 drafts', '−9% throughput; 99.2% cache occupancy at c64'],
    ['FP8 KV', 'More capacity; −25% versus paired BF16 at c64'],
], [3.5, 8.5], row_h=.87, size=23)
panel(s, 'Decide using latency, throughput, memory, and quality.')

s = base('Backup · Hard questions, short answers', 'Expanded answers and evidence: report.md §9')
table(s, ['Question', 'Answer'], [
    ['Error bars?', 'Most cells have one retained run; repeats are next.'],
    ['Why not fewer GPUs?', 'No sweep establishes the best feasible layout.'],
    ['Production cost?', 'Illustrative allocation cost; no quality or idle time.'],
    ['Real traffic replay?', 'No; workload summaries informed a synthetic grid.'],
    ['What could change the result?', 'Matched reruns or operator-level profiling.'],
], [3.6, 8.4], row_h=.68, size=21)

s = base('Backup · A CUDA error can start on the CPU', 'Source: fix_bug.md Bugs 2–3')
text(s, 'CUDA failure → compiler error → cache quota', .65, 1.98, 12, .7, 28, True, BLUE)
bullets(s, [
    'Disabling CUDA graphs did not resolve the failure.',
    'Eager execution still used the same compiler/cache path.',
    'Redirected caches allowed graph-enabled startup.',
], y=3.0, step=.8, size=24)
panel(s, 'A failed hypothesis test narrows the diagnosis.', size=24)

s = base('Backup · A failed route is not a failed idea', 'Source: fix_bug.md Bug 8; GLM alternate-stack controls', 'Experimental version-check bypass; numerical parity and production suitability unverified.')
table(s, ['GLM FP8 investigation', 'What it establishes'], [
    ['Layout needs 64 positional dims;\nGLM has 0', 'Selected route is incompatible'],
    ['Alternate NoPE path runs', 'FP8 KV is not universally impossible'],
    ['Paired BF16 348.8 → FP8 262.0 tok/s', '25% throughput tradeoff on that stack'],
], [6.3, 5.7], row_h=.91, size=21)
panel(s, 'Separate geometry, layout, backend, dtype, and version.', size=23)

assert len(prs.slides) == 20
for i, slide in enumerate(prs.slides, 1):
    assert slide.notes_slide.notes_text_frame.text.strip()
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, (i, shape.name)
        assert shape.left + shape.width <= prs.slide_width + 10, (i, shape.name)
        assert shape.top + shape.height <= prs.slide_height + 10, (i, shape.name)
prs.save(OUT)
check = Presentation(OUT)
assert len(check.slides) == 20
assert sum(s._element.get('show') == '0' for s in check.slides) == 8
print(f'Created {OUT}: 12 main slides, 8 hidden backup slides, 20 speaker notes.')
