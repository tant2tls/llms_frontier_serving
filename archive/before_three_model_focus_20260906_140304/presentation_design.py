"""Editable presentation visuals. Exact chart inputs are retained in an evidence ledger."""
from pathlib import Path
import json, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData

ROOT=Path(__file__).resolve().parents[1]
NAVY,INK,PAPER,WHITE='101F33','172C42','F6F7F9','FFFFFF'
MUTED,GRID,TEAL,PURPLE,ORANGE='586A7C','DCE3E9','007F78','7251B5','BB571E'
COLORS={'Qwen':TEAL,'GLM':PURPLE,'DeepSeek':ORANGE}
ARMS={'GLM':'GLM-5.3-Flash/results/bf16kv','DeepSeek':'deepseek_v4_flash/results/mtp-off-image','Qwen':'Qwen3.8-Flash-Next-FP8/results/base-util082'}
prs=Presentation()
prs.slide_width,prs.slide_height=Inches(13.333333),Inches(7.5)
prs.core_properties.title='Frontier MoE serving: the operating point matters'
prs.core_properties.author='Tan Ngo'
story,evidence=[],[]
MAIN_COUNT=16
TOTAL_COUNT=24
notes=['']*TOTAL_COUNT
REFERENCES={
    3:['https://arxiv.org/html/2606.19348v1','https://arxiv.org/html/2608.30320v1'],
    5:['https://kellerjordan.github.io/posts/muon/','https://arxiv.org/abs/1711.05101','https://github.com/KellerJordan/Muon','https://moonshotai.github.io/Kimi-K2/'],
    6:['https://github.com/deepseek-ai/DeepSeek-V3','https://arxiv.org/abs/2211.17192','https://www.lmsys.org/blog/2026-08-26-qwen-flash-next/'],
}

def rgb(c): return RGBColor.from_string(c)

def rect(s,x,y,w,h,fill):
    a=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    a.adjustments[0]=.08
    a.fill.solid(); a.fill.fore_color.rgb=rgb(fill); a.line.fill.background()
    return a

def text(s,value,x,y,w,h,size=22,bold=False,color=INK,align=None,record=True):
    box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,v in enumerate(str(value).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=v; p.font.name='Aptos'; p.font.size=Pt(size)
        p.font.bold=bold; p.font.color.rgb=rgb(color); p.space_after=Pt(5)
        if align is not None: p.alignment=align
    if record: story[-1]['text'].append(str(value))
    return box

def line(s,x1,y1,x2,y2,color=GRID,width=1):
    a=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    a.line.color.rgb=rgb(color); a.line.width=Pt(width)

def base(title,section,subtitle='',source='',caveat='',dark=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]); i=len(prs.slides)
    story.append(dict(title=title,text=[],tables=[],source=source,caveat=caveat))
    s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(NAVY if dark else PAPER)
    text(s,'UW SyFI  /  SERVING SYSTEMS',.55,.3,7,.23,10,True,'8DC9C3' if dark else TEAL,record=False)
    part='I / ' if 2<=i<=7 else 'II / ' if 8<=i<=MAIN_COUNT else ''
    text(s,(part+section).upper(),8.2,.3,4.55,.23,10,True,'A7BACB' if dark else MUTED,PP_ALIGN.RIGHT,False)
    text(s,title,.55,.85,12.1,.95,31,True,WHITE if dark else INK,record=False)
    if subtitle: text(s,subtitle,.58,1.77,12,.58,17,False,'BCD0DF' if dark else MUTED)
    if caveat: text(s,caveat,.6,6.57,12,.38,12,False,'BCD0DF' if dark else MUTED,record=False)
    line(s,.55,7.02,12.76,7.02,'34475A' if dark else GRID,.7)
    source_box=text(s,source or 'Tan Ngo · Professor Kan Zhu and UW SyFI · 6 September 2026',.58,7.12,11.5,.19,9,False,'BCD0DF' if dark else MUTED,record=False)
    if i in REFERENCES:
        source_box.text_frame.paragraphs[0].runs[0].hyperlink.address=REFERENCES[i][0]
    text(s,f'{i:02}' if i<=MAIN_COUNT else f'B{i-MAIN_COUNT}',12.05,7.09,.65,.25,11,True,'BCD0DF' if dark else MUTED,PP_ALIGN.RIGHT,False)
    s.notes_slide.notes_text_frame.text=notes[i-1].strip()
    if i>MAIN_COUNT: s._element.set('show','0')
    return s

def takeaway(s,value,dark=False):
    rect(s,.55,5.88,12.2,.53,'203A4C' if dark else 'E7F2EF')
    text(s,value,.76,5.99,11.75,.33,18,True,'B8E8DD' if dark else TEAL)

def card(s,x,y,w,h,kicker,headline,body,color=TEAL,dark=False):
    rect(s,x,y,w,h,'1C3045' if dark else WHITE)
    line(s,x+.02,y+.15,x+.02,y+h-.15,color,3)
    text(s,kicker.upper(),x+.2,y+.18,w-.4,.35,12,True,'A7D8D0' if dark else color)
    text(s,headline,x+.2,y+.68,w-.4,.90,24,True,WHITE if dark else INK)
    text(s,body,x+.2,y+1.55,w-.4,h-1.67,16,False,'C4D1DE' if dark else MUTED)

def table(s,headers,rows,widths,y=2.42,row_h=.62,size=18):
    t=s.shapes.add_table(len(rows)+1,len(headers),Inches(.6),Inches(y),Inches(12.1),Inches(row_h*(len(rows)+1))).table
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for r,row in enumerate([headers]+rows):
        for c,v in enumerate(row):
            a=t.cell(r,c); a.text=str(v)
            a.margin_left=Inches(.16); a.margin_right=Inches(.12)
            a.margin_top=a.margin_bottom=Inches(.05); a.vertical_anchor=MSO_ANCHOR.MIDDLE
            a.fill.solid(); a.fill.fore_color.rgb=rgb(NAVY if r==0 else WHITE if r%2 else 'EAF0F4')
            for p in a.text_frame.paragraphs:
                p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=r==0
                p.font.color.rgb=rgb(WHITE if r==0 else INK)
    story[-1]['tables'].append(dict(headers=headers,rows=rows))

def read(arm,file):
    path=f'{arm}/{file}.json'; d=json.loads((ROOT/path).read_text(encoding='utf-8'))
    assert d['completed']==d['num_prompts'] and d['failed']==0,path
    evidence.append(dict(slide=len(prs.slides),path=path,recorded_date=d.get('date'),output_throughput=d['output_throughput'],median_tpot_ms=d['median_tpot_ms'],median_ttft_ms=d['median_ttft_ms']))
    return d

def chart(s,cats,series,x,y,w,h,kind=XL_CHART_TYPE.LINE_MARKERS,ymax=None,ymin=0,unit=None,fmt='0',labels=False):
    data=CategoryChartData(); data.categories=list(map(str,cats))
    for name,values,_ in series: data.add_series(name,values)
    c=s.shapes.add_chart(kind,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
    c.font.name='Aptos'; c.font.size=Pt(14); c.font.color.rgb=rgb(MUTED)
    c.has_legend=len(series)>1
    if c.has_legend:
        c.legend.position=XL_LEGEND_POSITION.BOTTOM; c.legend.include_in_layout=False; c.legend.font.size=Pt(14)
    c.value_axis.minimum_scale=ymin
    if ymax is not None: c.value_axis.maximum_scale=ymax
    if unit is not None: c.value_axis.major_unit=unit
    c.value_axis.tick_labels.number_format=fmt; c.value_axis.has_major_gridlines=True
    c.value_axis.major_gridlines.format.line.color.rgb=rgb(GRID)
    for ax in (c.category_axis,c.value_axis):
        ax.tick_labels.font.size=Pt(14); ax.format.line.color.rgb=rgb(GRID)
    for ser,(_,_,col) in zip(c.series,series):
        ser.format.line.color.rgb=rgb(col); ser.format.line.width=Pt(2.8)
        ser.format.fill.solid(); ser.format.fill.fore_color.rgb=rgb(col)
        if kind==XL_CHART_TYPE.LINE_MARKERS:
            ser.marker.style=XL_MARKER_STYLE.CIRCLE; ser.marker.size=7
            ser.marker.format.fill.solid(); ser.marker.format.fill.fore_color.rgb=rgb(col)
            ser.marker.format.line.color.rgb=rgb(col)
    if labels:
        c.plots[0].has_data_labels=True; dl=c.plots[0].data_labels
        dl.position=XL_LABEL_POSITION.OUTSIDE_END; dl.font.size=Pt(16)
        dl.font.color.rgb=rgb(INK); dl.number_format=fmt
    story[-1]['tables'].append(dict(headers=['Category']+[z[0] for z in series],rows=[[cat]+[f'{z[1][j]:.3f}' for z in series] for j,cat in enumerate(cats)]))
    return c

s=base('Frontier MoE serving: the operating point matters','The thesis','GLM-5.3-Flash · DeepSeek-V4-Flash · Qwen3.8-Flash-Next-FP8',dark=True)
text(s,'Throughput, latency\nand memory lead to\ndifferent serving choices.',.65,2.58,7.05,2.02,31,True,WHITE)
for j,(v,l) in enumerate([('517.7','Qwen output tok/s · 16K, c64'),('12.89 s','GLM median TTFT · 131K, c8'),('+24%','GLM MTP n1 throughput · 16K, c1')]):
    yy=2.48+j*1.03
    text(s,v,8.05,yy,4.3,.54,30,True,'91D9C9'); text(s,l,8.07,yy+.58,4.3,.28,13,False,'C4D1DE')
takeaway(s,'8 × H100 80GB  |  Three measured deployments  |  Architecture + runtime + workload',True)

s=base('Two parts: architecture, then evidence','Research map','Research question: when does reduced model work become useful serving performance?',source='Source: task.md; report.md §§1–8')
for x,y,k,h,b,c in [(.6,2.48,'PART I / FRONTIER DESIGN','Where efficiency comes from','Sparse attention, hybrid state, MoE, Muon and native MTP.',TEAL),(6.8,2.48,'PART II / LOCAL MEASUREMENTS','Where the savings survive','Concurrency, context, prefix sharing, speculation and cost.',PURPLE),(.6,4.08,'METHOD / RESEARCH REASONING','Prediction → test → correction','Separate published mechanisms, local observations and hypotheses.',ORANGE),(6.8,4.08,'OUTCOME / A TESTABLE AGENDA','Optimize useful completions','Match runtime state; explain phase costs; test a serving policy.',TEAL)]:
    rect(s,x,y,5.9,1.42,WHITE)
    text(s,k,x+.2,y+.12,5.5,.24,11,True,c); text(s,h,x+.2,y+.44,5.5,.4,23,True)
    text(s,b,x+.2,y+.91,5.45,.46,14,False,MUTED)
takeaway(s,'Kimi K3 is architecture-only; no local performance or cost is extrapolated.')

s=base('Sparse attention reduces reads; selection has a cost','Sparse attention','A query needs useful history; the system must find, gather and process it.',source='Sources: DeepSeek-V4 report §2.3; Qwen3.8-Next report §2.1.2; report.md §2',caveat='Conceptual operator diagram. Sparse kernel savings and vendor comparisons are not local end-to-end speedups.')
for i,(title,kind,desc,col) in enumerate([('Dense history','dense','All available\npositions\n\nCore prefill:\nO(S²)',MUTED),('Sparse selection','sparse','Selected entries\nplus index cost\n\nCore attention:\nO(SK)',TEAL),('Compressed history','compressed','Summaries\nplus local tail\n\nFewer entries;\nextra state',PURPLE)]):
    xx=.6+i*4.12
    rect(s,xx,2.5,3.87,2.94,WHITE)
    text(s,title,xx+.18,2.69,3.51,.42,21,True,col)
    # Original schematic, not a copied research figure or measured attention pattern.
    # Rows are query positions; columns are history entries, with causal availability.
    count=10 if kind!='compressed' else 5
    cell=.153 if count==10 else .306
    for row in range(10):
        for c in range(count):
            causal=c<=row if count==10 else 2*c<=row
            selected=causal and (kind!='sparse' or c in {row,max(0,row-1),max(0,row-4)})
            a=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(xx+.18+c*cell),Inches(3.36+row*.153),Inches(cell-.025),Inches(.128))
            a.fill.solid(); a.fill.fore_color.rgb=rgb(col if selected else 'E7EDF1'); a.line.fill.background()
    text(s,'query × history',xx+.18,5.01,1.66,.25,10,False,MUTED)
    text(s,desc,xx+1.96,3.35,1.72,1.84,14,False,MUTED)
text(s,'Sparse path = indexer + top-k + gather + selected attention + state management',.74,5.5,12,.3,17,True,ORANGE)
takeaway(s,'Prediction: the crossover depends on context length, selection overhead, kernel efficiency and quality.')

s=base('Sparse expert compute; different memory obligations','Architecture','MoE selects a few experts per token, while the deployment retains a large weight footprint.',source='Source: report.md §2; recorded model analyses/configurations',caveat='Layer-count segments are schematic by type, not execution order or proportions of runtime.')
text(s,'ATTENTION LAYER COMPOSITION',3,2.43,5.9,.27,11,True,MUTED); text(s,'EXPERTS / SELECTED',10.32,2.43,2.3,.27,11,True,MUTED)
for i,(name,a,b,label,e) in enumerate([('GLM',34,11,'11 sparse','288 / 8'),('DeepSeek',0,43,'43 sparse · compressed','256 / 6'),('Qwen',36,12,'12 QSA','512 / 10'),('Kimi K3*',69,24,'24 gated MLA','896 / 16')]):
    yy=2.91+i*.65; wa=6.6*a/(a+b)
    text(s,name,.66,yy,2.25,.38,21,True)
    if a:
        rect(s,3,yy,wa,.43,TEAL); text(s,f'{a} recurrent',3.1,yy+.08,wa-.13,.27,13,True,WHITE)
    rect(s,3+wa,yy,6.6-wa,.43,PURPLE); text(s,label,3+wa+.08,yy+.08,6.44-wa,.27,11,True,WHITE)
    text(s,e,10.35,yy+.01,2.15,.38,20,True)
text(s,'Recurrent: fixed state / active sequence',.66,5.62,5.9,.25,13,True,TEAL)
text(s,'Retained history: grows with sequence length',6.65,5.62,6,.25,13,True,PURPLE)
takeaway(s,'Budget weights + workspace + graphs + history + recurrent state.   *Kimi unmeasured')

s=base('Muon changes training updates; AdamW still has a role','Training efficiency','A matrix-aware optimizer is a training design choice, not a decode kernel.',source='Sources: Jordan, Muon; Loshchilov & Hutter, AdamW; Qwen/DeepSeek reports; Kimi K2 report',caveat='Published optimizer recipes; no local training comparison. Do not attribute measured inference throughput to Muon.')
card(s,.6,2.45,5.92,2.97,'ADAMW','Coordinate-wise scaling','Momentum + second-moment scaling\nDecoupled weight decay\nEmbeddings and other selected groups',MUTED)
card(s,6.76,2.45,5.92,2.97,'MUON','Matrix update geometry','Momentum → approximate orthogonalization\nNewton–Schulz matrix operations\nMatrix-aware partitioning and batching',PURPLE)
text(s,'Qwen and DeepSeek use mixed recipes; Kimi K2 reports MuonClip for training stability.',.73,5.48,12,.31,17,True)
takeaway(s,'Evaluate quality reached per training GPU-hour; the optimizer update is absent from serving.')

s=base('Day-0 speculation: ship a draft, still pay for verification','Native MTP','Here “zero-day” means native drafting and runtime support at release, not zero overhead.',source='Sources: DeepSeek-V3 MTP; Leviathan et al., speculative decoding; SGLang Qwen day-0 support',caveat='Runtime support is version-specific. Target-distribution preservation requires correct acceptance/resampling and state handling.')
for i,(k,h,b,col) in enumerate([('01 / DRAFT','Propose future tokens','Native MTP can avoid waiting for a separately trained draft model.',TEAL),('02 / VERIFY','Target checks candidates','Batch candidate verification; keep valid progress and resample as required.',PURPLE),('03 / COMMIT','Restore consistent state','Discard rejected suffix state; commit KV, indices and recurrent updates.',ORANGE)]):
    card(s,.6+i*4.12,2.5,3.87,2.88,k,h,b,col)
text(s,'Break-even:  round time / expected committed tokens  <  ordinary time / token',.72,5.5,12,.3,17,True,TEAL)
takeaway(s,'Availability does not guarantee acceleration: acceptance, draft cost, verification and load set the payoff.')

s=base('An architectural saving must survive the whole request','From design to hypotheses','Analytical lens: phase-local improvements compete with input work, state, communication and scheduling.',source='Source: report.md §2 analytical framework; §8 proposed experiments',caveat='Illustrative Amdahl calculation, not a measured time breakdown. No dense control or local operator attribution is available.',dark=True)
text(s,'If attention is 30% of runtime,\na 10× attention speedup gives…',.7,2.57,7.2,1.1,28,True,WHITE)
text(s,'1 / (0.70 + 0.30 / 10) = 1.37×',.73,3.97,7.35,.65,27,True,'91D9C9')
card(s,8.55,2.49,4.1,3.08,'TESTABLE PREDICTIONS','Measure what moved','Sparse state → context scaling\nMoE → load / communication\nMTP → accepted progress / round\nPool sizing → capacity / preemption',TEAL,True)
takeaway(s,'Part II tests the operating regimes; a causal speedup claim still needs a matched intervention.',True)

s=base('A fixed GPU budget; three workload axes','Experimental design','8 × H100 80GB · TP8 + expert parallelism · 256 output tokens · synthetic text prompts',source='Source: bench.sh; experiments.md; report.md §1',caveat='Mostly one retained run per point. GPU class/count match; builds, backends and some startup pools differ.')
for i,(a,b,c) in enumerate([('BATCH','16K input','c = 1, 4, 16, 64\n8, 8, 32, 128 requests'),('CONTEXT','16K → 260K input','c = 8\n16 requests per point'),('PREFIX','64K shared + 2K suffix','1, 4, 16 distinct prefixes\nc = 8 · 64 requests')]):
    card(s,.6+i*4.12,2.47,3.87,2.7,a,b,c,[TEAL,PURPLE,ORANGE][i])
text(s,'ARM BOUNDARY',.65,5.38,1.8,.25,11,True,ORANGE)
text(s,'Qwen batch: base-util082   |   Qwen context/prefix: earlier base pool',2.44,5.32,10.1,.4,17,True)
takeaway(s,'Concurrency caps client requests; output tok/s includes input work; TTFT includes queueing.')

s=base('Qwen leads the corrected 16K throughput grid','01 / Throughput','Output tokens/s ↑ · 16,384 input / 256 output · MTP off',source='Source: main batch JSONs — Qwen base-util082; GLM bf16kv; DeepSeek mtp-off-image',caveat='Deployment comparison; quality unmeasured. Lines connect measured concurrency categories, not evenly spaced numeric intervals.')
cs=[1,4,16,64]; batch={m:[read(a,f'batch_isl16k_c{c}') for c in cs] for m,a in ARMS.items()}
chart(s,cs,[(m,[d['output_throughput'] for d in batch[m]],COLORS[m]) for m in ['Qwen','GLM','DeepSeek']],.55,2.4,8.5,3.27,ymax=600,unit=200)
text(s,'Client concurrency cap',3,5.62,4,.23,12,False,MUTED)
card(s,9.36,2.54,3.3,2.97,'At concurrency 64','517.7 tok/s','Qwen / GLM       1.16×\nQwen / DeepSeek 1.33×\n\nObserved throughput lead',TEAL)
takeaway(s,'64× more concurrency yields only 4.6–4.9× more output throughput.')

s=base('Past c8, latency grows much faster than throughput','01 / Operating point','Each metric normalized to its own c8 value = 1.0× · median TPOT = time per output token',source='Source: GLM bf16kv; DeepSeek util085-dev20073 finer batch grids',caveat='DeepSeek is a separate build/utilization arm. c8 is an anchor, not a proven optimum. Horizontal positions are concurrency categories.')
for j,(m,a) in enumerate([('GLM',ARMS['GLM']),('DeepSeek','deepseek_v4_flash/results/util085-dev20073')]):
    ds=[read(a,f'batch_isl16k_c{c}') for c in [8,16,32,48,64]]
    t=[d['output_throughput']/ds[0]['output_throughput'] for d in ds]; p=[d['median_tpot_ms']/ds[0]['median_tpot_ms'] for d in ds]
    text(s,f'{m}: {t[-1]:.2f}× throughput / {p[-1]:.2f}× TPOT',.7+j*6.22,2.43,6.05,.4,18,True,COLORS[m])
    chart(s,[8,16,32,48,64],[('Throughput',t,TEAL),('Median TPOT',p,ORANGE)],.55+j*6.22,2.97,6,2.7,ymax=10,unit=2,fmt='0.0"×"')
takeaway(s,'Choose concurrency against a latency objective; aggregate token rate alone cannot choose it.')

s=base('Long input changes which metric wins','01 / Context','131,072 input / 256 output · c8 · 16 requests',source='Source: ctx_isl131072_c8.json — GLM bf16kv; DeepSeek mtp-off-image; Qwen base',caveat='Qwen uses the earlier smaller pool. Rerun matched startup before interpreting long-context capacity as architectural.')
ctx={m:read(a if m!='Qwen' else 'Qwen3.8-Flash-Next-FP8/results/base','ctx_isl131072_c8') for m,a in ARMS.items()}
for j,(field,title,div,maxv) in enumerate([('output_throughput','Output tokens/s ↑',1,70),('median_ttft_ms','Median first-token latency, seconds ↓',1000,34)]):
    text(s,title,.7+j*6.22,2.46,6,.4,20,True)
    c=chart(s,['Qwen','GLM','DeepSeek'],[('Observed',[ctx[m][field]/div for m in ['Qwen','GLM','DeepSeek']],TEAL)],.6+j*6.22,3.02,5.95,2.58,kind=XL_CHART_TYPE.COLUMN_CLUSTERED,ymax=maxv,fmt='0.0' if j==0 else '0.00',labels=True)
    for point,m in zip(c.series[0].points,['Qwen','GLM','DeepSeek']):
        point.format.fill.solid(); point.format.fill.fore_color.rgb=rgb(COLORS[m]); point.format.line.fill.background()
takeaway(s,'Qwen has the highest observed output rate; GLM has the lowest observed median TTFT.')

s=base('Prefix reuse is a state-consistency problem','02 / Prefix caching','64 requests · 64K shared + 2K unique · c8 · prefix caching enabled throughout',source='Source: GLM/DeepSeek prefix_p64k_n*.json; report.md §4',caveat='Sharing patterns, not cache-on/off speedups. Fill order, eviction and scheduling explanations remain hypotheses.')
text(s,'Output tokens/s ↑',.7,2.43,6,.35,18,True)
chart(s,[1,4,16],[(m,[read(ARMS[m],f'prefix_p64k_n{n}')['output_throughput'] for n in [1,4,16]],COLORS[m]) for m in ['GLM','DeepSeek']],.55,2.9,6.35,2.72,ymax=450,unit=150)
text(s,'Distinct shared prefixes (categories)',1.8,5.63,4.8,.23,12,False,MUTED)
rect(s,7.28,2.51,5.35,3.1,WHITE); text(s,'Restore the same prefix boundary',7.5,2.73,4.92,.42,20,True)
for yy,label,col in [(3.42,'Retained KV / compressed history',PURPLE),(4.03,'Recurrent-state checkpoint',TEAL)]:
    rect(s,7.53,yy,4.65,.43,col); text(s,label,7.7,yy+.07,4.3,.27,14,True,WHITE)
line(s,11.9,3.2,11.9,4.65,ORANGE,2)
text(s,'Then branch safely; copy or restore mutable state.',7.51,4.82,4.75,.57,16,False,MUTED)
takeaway(s,'GLM and DeepSeek peak at four prefixes; maximum sharing did not maximize throughput.')

s=base('Speculation helps at light load; the gain disappears','03 / MTP','GLM throughput with MTP ÷ base throughput · 16K input / 256 output',source='Source: GLM bf16kv / bf16kv-mtp-n1 / bf16kv-mtp-n5; paired context arms',caveat='GLM supplies the strongest controls. Real-text acceptance and quality are untested; scheduling attribution remains open.')
ratios=[]
for n,col in [(1,TEAL),(5,PURPLE)]:
    ratios.append((f'{n} draft token'+('s' if n>1 else ''),[read(f'GLM-5.3-Flash/results/bf16kv-mtp-n{n}',f'batch_isl16k_c{c}')['output_throughput']/batch['GLM'][j]['output_throughput'] for j,c in enumerate(cs)],col))
chart(s,cs,ratios+[('Base = 1×',[1]*4,MUTED)],.55,2.45,8.4,3.2,ymin=.8,ymax=1.3,unit=.1,fmt='0.0"×"')
card(s,9.26,2.54,3.38,2.98,'131K context / n1','−17% TTFT','But −4% output rate.\n\nLatency and allocation cost move in opposite directions.',PURPLE)
takeaway(s,'At c1: +24% with one draft. At c64: approximately flat with one draft; −9% with five.')

s=base('Lower token cost is purchased with higher latency','04 / Allocation cost','Illustrative $2.50/GPU-hour × 8 GPUs = $20/node-hour · 16K input / 256 output',source='Source: exact main batch JSONs; report.md §6. Cost = 20 × 1,000,000 / (3,600 × tok/s)',caveat='Includes input work. No quality adjustment, idle-time model, price quote or measured cost per successful task.')
text(s,'Dollars per million output tokens ↓',.7,2.44,7.6,.4,20,True)
chart(s,['Qwen','GLM','DeepSeek'],[(f'c{cs[j]}',[20e6/(3600*batch[m][j]['output_throughput']) for m in ['Qwen','GLM','DeepSeek']],col) for j,col in [(0,'8698A9'),(3,TEAL)]],.55,2.98,8.55,2.72,kind=XL_CHART_TYPE.COLUMN_CLUSTERED,ymax=80,fmt='$0.00',labels=True)
card(s,9.33,2.56,3.3,3,'GLM / c1 → c64','$57.68 → $12.43','Cost / million outputs\n\nMedian TPOT:\n7.1 → 129.5 ms',PURPLE)
takeaway(s,'A serving objective needs useful completions within latency targets; token cost is only one input.')

s=base('Startup state can impersonate an architecture effect','Debugging / Bug 12','Qwen c4: 162.2 tok/s in the earlier arm → 256.8 tok/s in the corrected batch arm.',source='Source: Qwen base-util082/RESULT-util-ab.md, both c4 JSONs; fix_bug.md Bug 12',caveat='Startup and utilization changed together. Preemption causality is unmeasured; the batch correction does not transfer to other axes.')
for i,(k,h,b,c) in enumerate([('STARTUP ACCOUNTING','17.07 → 0.99 GiB','Recorded peak activation\nCold → warm compile cache',ORANGE),('REPORTED CACHE POOL','2.05M → 3.20M','Token capacity\nUtilization 0.85 → 0.82',PURPLE),('OBSERVED c4 RESULT','162.2 → 256.8','Output tokens/s\nCheckpoint unchanged',TEAL)]):
    card(s,.6+i*4.12,2.5,3.87,2.86,k,h,b,c)
takeaway(s,'Inspect runtime state before assigning an architectural cause to a changed ranking.')

s=base('Research direction: budget state and speculation together','Conclusions / Discussion','Hypothesis: load-aware draft budgets can improve goodput when hybrid-state capacity is constrained.',source='Source: report.md §8; fix_bug.md §6. Research hypothesis and evaluation below are proposed.',dark=True)
for i,(k,h,b) in enumerate([('MY EVIDENCE','Operating points matter','GLM MTP helps at light load.\nQwen startup changes the ranking.\nLatency and token cost diverge.'),('FIRST / ESTABLISH','Match and profile','Repair gates; repeat matched arms.\nSeparate prefill / decode costs.\nMeasure acceptance and state work.'),('THEN / TEST','Adaptive draft budget','Compare off, fixed, adaptive.\nUseful completions under SLOs.\nQuality, overhead and tails.')]):
    card(s,.6+i*4.12,2.48,3.87,3,k,h,b,[TEAL,PURPLE,ORANGE][i],True)
takeaway(s,'Falsifier: no useful gain over fixed policies after matching quality, latency targets and runtime state.',True)

s=base('What is controlled, and what remains different?','Backup / Comparability',source='Source: report.md §1; experiments.md')
table(s,['Held or documented','Different / missing'],[['8 × H100 80GB; TP8 + EP','Physical node/topology equivalence unestablished'],['Synthetic input/output length targets','Tokenizers, actual counts and task quality'],['MTP off in main comparison','Weight/KV precision, builds and backends'],['Named arm and per-point evidence','Some startup pools; session/repeat coverage']],[4.4,7.7],row_h=.67)
takeaway(s,'DeepSeek’s build bridge applies to its batch grid only; all models are not engine-matched.')

s=base('A workload curve does not identify a dominant kernel','Backup / Bottleneck attribution',source='Source: report.md §3; fix_bug.md Bug 13',caveat='GLM engine estimates omit attention. Whole-request averages can hide bandwidth-bound individual kernels.')
table(s,['Candidate mechanism','Discriminating evidence needed'],[['Small GEMMs / expert routing','Operator durations, shapes and per-rank expert load'],['Collectives / launch overhead','Timeline and controlled TP/EP or graph intervention'],['HBM traffic','Hardware counters within prefill/decode intervals'],['Scheduler / cache pressure','Per-step batches, preemption and matched-pool controls']],[4.4,7.7],row_h=.67)
takeaway(s,'Enabled → implemented → emitted → scraped → aligned → saved: audit the metric chain.')

s=base('A reusable prefix needs a complete state checkpoint','Backup / Hybrid cache',source='Source: report.md §4; fix_bug.md Bug 6')
for i,(k,h,b) in enumerate([('IDENTITY','Same token prefix','Model revision, adapter, positions and relevant execution state must match.'),('BOUNDARY','All layers agree','KV/history and recurrent state resume at the same compatible position.'),('BRANCH / ROLLBACK','Protect mutable state','Share immutable history; copy or restore state when requests diverge.')]):
    card(s,.6+i*4.12,2.46,3.87,3.05,k,h,b,[TEAL,PURPLE,ORANGE][i])
takeaway(s,'Finer checkpoints trade additional memory/copy work for less recomputation after partial matches.')

s=base('Resident parameters are not per-token arithmetic','Backup / Weight accounting',source='Source: report.md §2; prior tensor analyses; recorded Kimi official-card reference',caveat='Historical tensor analyses, not a new recount. Weight payload lower bounds are not validated deployments.')
table(s,['Model','Total / active GEMM parameters','Additional obligation'],[['GLM','321.34B / 17.38B','Recurrent state + retained history'],['DeepSeek','290.91B / 14.08B','Compression and sparse indexing'],['Qwen','176.94B served / 7.27B','51.23B n-gram lookup table']],[2.3,5,4.8],row_h=.71)
text(s,'Kimi K3: 2.8T × 0.5 bytes ≈ 1,304 GiB of ideal four-bit weights',.7,5.51,12,.34,18,True,PURPLE)
takeaway(s,'MoE reduces selected compute; weights, communication, lookup and state still need accounting.')

s=base('More capacity and more drafts are not automatic wins','Backup / MTP and KV',source='Source: GLM paired arms; report.md §§5, 7',caveat='Alternate FP8 stack bypasses a version check; numerical parity and production suitability remain unverified.')
table(s,['GLM option','Measured tradeoff','Decision supported'],[['MTP n1 / c1','+24% throughput; TPOT 7.06 → 5.05 ms','Candidate for light-load testing'],['MTP n5 / c64','−9% throughput; 99.2% cache occupancy','More drafts do not ensure benefit'],['Alternate FP8 KV / c64','−25% versus paired alternate BF16','Price capacity against speed']],[2.5,5.7,3.9],row_h=.88,size=17)
takeaway(s,'Reported FP8 capacity gain = 1.805× in its startup comparison; preserve session boundaries.')

s=base('Answer directly, then name the missing experiment','Backup / Defense',source='Expanded answers: report.md §9; fix_bug.md §7')
table(s,['Question','Defensible answer'],[['Where are error bars?','Most cells have one retained run; randomized repeats are pending.'],['Cheapest in production?','Unknown: no quality, real arrival trace, idle-time or GPU-count sweep.'],['How fast is prefix caching?','Only sharing patterns were measured; a cache-off control is missing.'],['Why not name the bottleneck?','Request metrics do not isolate operators; capture phase-specific traces.'],['What could change the ranking?','Matched startup/pools, repeats, different workloads and quality.']],[3.35,8.75],row_h=.59,size=17)

s=base('Read the earliest cause, not the loudest CUDA symptom','Backup / Bugs 2–3',source='Source: fix_bug.md Bugs 2–3; GLM cu130 run6_eager and run7_trtllmfix logs',caveat='Logs support the failure chain; they do not establish every internal kernel-handle transition.')
for i,(k,h,b) in enumerate([('SYMPTOM','CUDA invalid argument','Eager execution did not resolve startup; inspect shared dependencies.'),('EARLIER EVIDENCE','Cache write fails','TensorRT/DeepGEMM preparation hits a separate filesystem path.'),('OBSERVED RESOLUTION','Redirect the caches','Redirect backend-specific caches and add write preflight; graph-enabled startup succeeds.')]):
    card(s,.6+i*4.12,2.5,3.87,3,k,h,b,[ORANGE,PURPLE,TEAL][i])
takeaway(s,'A failed intervention narrows the cause; eager execution can still use the same compiler cache.')

s=base('Separate the backend change from the dtype change','Backup / Bug 8',source='Source: GLM c64 JSONs — bf16kv, bf16kv-fi618, fp8kv-fi618; fix_bug.md Bug 8',caveat='Failed layout required 64 positional dimensions; GLM NoPE has zero. Alternate execution does not validate numerical parity.')
fp=[read(f'GLM-5.3-Flash/results/{a}','batch_isl16k_c64')['output_throughput'] for a in ['bf16kv','bf16kv-fi618','fp8kv-fi618']]
text(s,'Output tokens/s ↑ · concurrency 64',.7,2.44,7.8,.4,20,True)
chart(s,['Original BF16','Alternate BF16','Alternate FP8'],[('GLM',fp,PURPLE)],.6,2.96,8.5,2.68,kind=XL_CHART_TYPE.COLUMN_CLUSTERED,ymax=550,fmt='0.0',labels=True)
card(s,9.35,2.55,3.3,3,'PAIRED COMPARISONS','−22% then −25%','Combined: −41%, not −47%.\n\n0.780 × 0.751 ≈ 0.586\nof original throughput',ORANGE)
takeaway(s,'Geometry, layout, kernels and versions determine whether a dtype has a working execution path.')

def build():
    assert len(prs.slides)==TOTAL_COUNT
    # script.md is the complete spoken text; notes also retain timing and evidence cues.
    scripts={}
    script_path=ROOT/'script.md'
    if script_path.exists():
        for n,body in re.findall(r'<!-- SCRIPT (\d+) -->\s*(.*?)\s*<!-- END SCRIPT -->',script_path.read_text(encoding='utf-8'),re.S): scripts[int(n)]=body
        assert set(scripts)==set(range(1,TOTAL_COUNT+1))
    timing=re.findall(r'\*\*(\d+:\d+) · cumulative (\d+:\d+)',script_path.read_text(encoding='utf-8'))
    assert len(timing)==MAIN_COUNT
    for i,s in enumerate(prs.slides,1):
        cue=f'SLIDE {i} — {timing[i-1][0]}; cumulative {timing[i-1][1]}.' if i<=MAIN_COUNT else f'BACKUP B{i-MAIN_COUNT} — use only for relevant questions.'
        notes[i-1]=cue+('\n\nFULL SPOKEN SCRIPT\n'+scripts[i] if scripts else '')
        if i in REFERENCES:
            notes[i-1]+='\n\nPRIMARY REFERENCES (checked 6 September 2026)\n'+'\n'.join(REFERENCES[i])
        s.notes_slide.notes_text_frame.text=notes[i-1]
        for sh in s.shapes:
            assert sh.left>=0 and sh.top>=0,(i,sh.name)
            assert sh.left+sh.width<=prs.slide_width+10,(i,sh.name)
            assert sh.top+sh.height<=prs.slide_height+10,(i,sh.name)
    out=ROOT/'SyFI_ML_Serving_refined.pptx'; prs.save(out)
    md=['---','marp: true','theme: default','paginate: true','size: 16:9','---','',
        '> Two parts: frontier architecture (slides 2–7), then measurements and research analysis (slides 8–16). 16 main slides (18 minutes) + 8 hidden backups. [Full speaking script](script.md). '
        'Editable charts and diagrams are in the PowerPoint. The builder synchronizes visible content; '
        'chart tables retain three decimals for review. Edit spoken text in script.md and rehearsal cues in the comments below.','']
    for i,st in enumerate(story):
        if i: md+=['---','']
        md+=[f"## {i+1 if i<MAIN_COUNT else 'Backup B'+str(i-MAIN_COUNT+1)}. {st['title']}",'']
        for tx in st['text']: md+=[tx.replace('\n','  \n'),'']
        for tab in st['tables']:
            clean=lambda v:str(v).replace('|','/').replace('\n','<br>')
            md+=['| '+' | '.join(map(clean,tab['headers']))+' |','| '+' | '.join(['---']*len(tab['headers']))+' |']
            md+=['| '+' | '.join(map(clean,r))+' |' for r in tab['rows']]; md+=['']
        if st['caveat']: md+=['**Boundary:** '+st['caveat'],'']
        md+=[st['source'] or 'Tan Ngo · Professor Kan Zhu and UW SyFI · 6 September 2026','','<!--',notes[i],'-->','']
    (ROOT/'slides.md').write_text('\n'.join(md),encoding='utf-8')
    (ROOT/'artifacts').mkdir(exist_ok=True)
    (ROOT/'artifacts/presentation_evidence.json').write_text(json.dumps(dict(scope='Exact JSON inputs for charts; formulas in presentation_design.py; other facts sourced in report.md.',points=evidence),indent=2),encoding='utf-8')
    assert sum(s._element.get('show')=='0' for s in prs.slides)==8
    print(f'Created {out}: {MAIN_COUNT} main + 8 hidden backups; {sum(sh.has_chart for s in prs.slides for sh in s.shapes)} editable charts; {TOTAL_COUNT} notes.')
    print('Synchronized slides.md and artifacts/presentation_evidence.json.')
